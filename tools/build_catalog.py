"""Genera catalog.js a partir del Excel de campaña.

Uso: python tools/build_catalog.py ruta/al/Cyberpunk.xlsx > catalog.js
Hojas leídas: Ojos, Corteza Frontal, Sistema operativo, Méritos.
Los méritos generales (imagen de la tabla) y los implantes que no están en el Excel van abajo, a mano.
"""
import json
import re
import sys

import openpyxl

SHEETS = {  # hoja -> (ranura, columna de precio, columna de icewall)
    'Ojos': ('Cara / Ojos', 'D', None),
    'Corteza Frontal': ('Corteza frontal', 'D', None),
    'Sistema operativo': ('Sistema operativo', 'E', 'D'),
}

# Implantes que no figuran en el Excel (cargados por el jugador en la versión anterior).
EXTRA_IMPLANTS = [
    {'name': 'Kerenzikov', 'slot': 'Sistema nervioso', 'minTier': 1, 'load': 12, 'load6': 18,
     'price': 9000, 'tiers': ['+2 PA, +10 Reflejos, D 1T, CD 3T', '+5 Reflejos', '+10 Evasión', '+1 PA', '+5 Crítico',
                              'Duplica los efectos anteriores']},
    {'name': 'Defenzikov', 'slot': 'Sistema integumentario', 'minTier': 1, 'load': 15, 'load6': 22,
     'price': 10000, 'tiers': ['+5 RD, D 1T, CD 3T. Se activa al terminar Kerenzikov', '+5 RD', '+5 RD', '+5 RD',
                               '+5 RD, D 2T', 'Inmune a críticos mientras Kerenzikov está activo']},
]

# Méritos generales (tabla de la imagen): nivel, nombre, requisito, rangos, descripción.
GENERAL_MERITS = """
1|Ambidiestro|DES 60|1|Reduce a la mitad los penalizadores por usar la mano torpe
1|Regenerador|CON 60|1|+2 Velocidad de curación
1|Entrenado|-|6|+5 Característica
1|Experto|Hab. 50|6|+20 a una habilidad especializada que no sea de combate (Ej: Investigar: Rastrear, Pericia: Medicina, Social: Comerciar)
1|Habilidoso con las armas|DES 60|1|+30 Fuerza para cumplir los requisitos mínimos de armas
1|Instruido|INT 40|1|+3 PHxNv
1|Manos de acero|Atletismo 40|1|No sufre penalizadores por escalar sin equipamiento
1|Parchar|Pericia 40|1|Reduce un 50% mas el sangrado al realizar primeros auxilios
1|Primera sangre|-|1|+3 Iniciativa
1|Uno con la naturaleza|-|1|Mitad de tiempo para cazar, encontrar agua y montar campamento
2|Agarre experto|Pelea: Agarres 40|1|+20 a fuerza o destreza al mantener una presa
2|Decidido|-|1|+10 Voluntad
2|El Alma de la Fiesta|-|1|+20 sociales bajo la influencia del alcohol
2|Escapista|-|1|+20 a destreza o agilidad para intentar soltarse de ataduras o agarres
2|Escurridizo|-|1|+10 Reflejos
2|Imparable|-|1|+10 Fortaleza
2|Incansable|-|1|+30 a fortaleza para resistir el cansancio
2|Sanador|Pericia: Medicina 50|1|Aumenta los puntos de vida curados con Pericia: Medicina en 50%
2|Siempre listo|DES 60|1|-2 PA para desenfundar/desenvainar un arma o tomar un objeto de acceso rápido
2|Sin tiempo para sangrar|-|1|Daño por desangre solo cada 2 turnos
3|Brazo de hierro|FUE 50|1|Duplica la potencia para arrojar
3|Brazo firme|FUE 70|1|Reduce a la mitad el penalizador de armas automáticas
3|Caer y rodar|Atletismo 60|1|Mitad de daño por caídas
3|Duro de matar|DET 50|1|+20 Vitalidad
3|Lider|Social 70|1|+2 PA, +5 evasión y +10 a todas las habilidades a los miembros del grupo en un rango de 10m
3|Observador|Investigación 70|1|Permite saber los pv, resistencias y armamento del enemigo
3|Ojo de halcón|PER 50|1|Duplica el rango de tiro
3|Médico de combate|Pericia: Medicina 60|1|Mitad de PA para curar en combate
4|Descarga de Adrenalina|DET 60|1|+20 Fuerza, +1 a los PA, +5 RD y reduce a la mitad los penalizadores cuando tus PGs bajen del 50%
4|Afilado|-|2|+5 crítico en combate
4|CQC|Puntería 50, Pelea 50|1|Elimina el penalizador por disparar a corto rango
4|Escurridizo II|DES 70|1|+1 evasión por cada metro recorrido
4|Rata de laboratorio|CON 50|1|Reduce los efectos negativos de sustancias consumidas en 50%
4|Soy tu Huckleberry|Puntería 60|1|Permite realizar disparos apuntados a 15m o menos por el coste de un disparo rápido
4|Visión nocturna|PER 60|1|Mitad de penalizador por baja luminosidad
4|Evasivo|-|2|+5 Evasión
5|99.1% puro|INT 60, CON 40|1|Aumenta los efectos positivos de sustancias consumidas en un 50%
5|Acróbata|Atletismo: Acrobacia 60|1|Mitad de PA para realizar maniobras acrobáticas
5|Dureza|CON 60, DET 60|2|+5% RD y +10 Fortaleza
5|Ejecutor|-|2|+1 resultado de crítico
5|Listo para la acción|-|2|+2 acción en el primer turno de combate, +1 en el resto
5|Recarga táctica|DES 40|1|Mitad de PA para acciones de recarga
6|Cuidadoso|-|1|Reduce la probabilidad de pifia en 3%
6|Gatillo rápido|DES 70|1|-1 PA para con armas de distancia
6|Guerrero del asfalto|Pericia: Conducir 60|1|Sin penalizador atacar en un vehículo en movimiento
6|Manos rápidas|DES 70|1|-1 PA para combate cuerpo a cuerpo
6|Experto en combate|Hab. 50|3|+20 a una habilidad especializada de combate (Ej: Pelea: Judo, Puntería: Rifles)
6|Puños de acero|Pelea 60|1|Reduce la RD enemiga a la mitad al atacar con las manos desnudas
7|Fundirse en las sombras|Sigilo 70|1|+5 al crítico para ocultarse en las sombras, +15 en la oscuridad completa. No hay pifia de sigilo
7|Esquivador de balas|Reflejos 40|1|Permite usar reflejos para esquivar ataques a distancia sin penalizador
7|Francotirador|Puntería 70, PER 60|1|Reduce en 20 el penalizador por apuntar
7|Cirujano de combate|Pericia: Medicina 70|1|Permite realizar cirugías complejas en combate
7|Sentido del Peligro|PER 75|1|Siempre que vayas a entrar en una situación peligrosa (lo sepas o no), se te erizará el pelo de la nuca y sentirás que algo malo va a suceder
7|Experto en Demolición|Pericia: Explosivos 70|1|+50% de daños de explosivos y siempre explotarán en el momento correcto
8|Paso silencioso|Sigilo 65, Atletismo 65|1|Puede mantener el sigilo sin penalizador a paso rápido
8|Peligro a dos manos|DES 65|1|Permite realizar dos ataques portando un arma de igual tamaño en cada mano por un 20% menos de PA
8|Tocado por la fortuna|-|1|Los usos de suerte cuestan la mitad
8|Fuerza de voluntad|DET 80|1|Los usos de resolución cuestan la mitad
8|Manos mágicas|Pericia: Medicina 80|1|La curación de Pericia: Medicina siempre es el máximo posible
9|Asesino|Sigilo 80, Pelea 80|1|+20 crítico, +2 letalidad al realizar un ataque cuerpo a cuerpo estando oculto
9|Gun kata|Pelea 70, Puntería 70, DES 70|1|Luego de una tirada exitosa de reflejos permite realizar un disparo apuntado por la mitad de PA
9|Rendirse, jamás|DET 80|1|La resolución no puede bajar de 5
9|Infalible|-|1|La tirada de efecto de pifias se considera siempre 1
9|Trébol de 4 hojas|-|1|La suerte no puede bajar de 5
10|El profesional|-|1|Duplica la probabilidad de crítico de todas las habilidades
10|Carrera de la Parca|-|1|Cada vez que elimines a un oponente recuperas automáticamente 3 Puntos de acción (no se puede exceder del máximo de PA por turno)
10|El destructor|-|1|Cada golpe CaC que realices es crítico con una tirada de (FUE+DES)/2
10|Una bala, una muerte|-|1|Cada ataque a distancia que realices es crítico con una tirada de (DES+PER)/2
10|Soy Leyenda|-|1|Los rumores de tus hazañas (positivas o negativas) son reconocidos mundialmente, muchas veces incluso exagerados
"""


def num(v):
    return v if isinstance(v, (int, float)) else None


def implants(wb):
    out = []
    for sheet, (slot, price_col, ice_col) in SHEETS.items():
        ws = wb[sheet]
        block = []
        for r in range(2, ws.max_row + 1):
            text = ws[f'C{r}'].value
            if text:
                block.append(r)
            name = ws[f'A{r}'].value
            if not name or not block:  # el nombre está en la última fila de cada bloque
                continue
            tiers, minTier = [None] * 6, 6
            for br in block:
                m = re.match(r'\s*TIER\s*(\d)\s*\|\s*(.*)', str(ws[f'C{br}'].value), re.S)
                if not m:
                    continue
                t, rest = int(m.group(1)), re.sub(r'\s+', ' ', m.group(2)).strip()
                if ice_col and num(ws[f'{ice_col}{br}'].value) is not None:
                    rest = f'{ws[f"{ice_col}{br}"].value} Icewall | {rest}'
                tiers[t - 1] = rest
                minTier = min(minTier, t)
            loads = [(br, num(ws[f'B{br}'].value)) for br in block if num(ws[f'B{br}'].value) is not None]
            load6 = next((v for br, v in loads if str(ws[f'C{br}'].value).strip().startswith('TIER 6')), None)
            load = next((v for br, v in loads if not str(ws[f'C{br}'].value).strip().startswith('TIER 6')), load6)
            price = next((num(ws[f'{price_col}{br}'].value) for br in block if num(ws[f'{price_col}{br}'].value)), None)
            out.append({'name': str(name).strip().title(), 'slot': slot, 'minTier': minTier,
                        'load': load or 0, 'load6': load6 or load or 0, 'price': price, 'tiers': tiers})
            block = []
    return out + EXTRA_IMPLANTS


def merits(wb):
    rows = [l.split('|') for l in GENERAL_MERITS.strip().splitlines()]
    out = [{'lvl': int(a), 'name': b, 'req': c, 'max': int(d), 'desc': e} for a, b, c, d, e in rows]
    ws = wb['Méritos']
    for r in range(2, ws.max_row + 1):
        lvl, name, req, ranks, desc = (ws.cell(r, c).value for c in range(1, 6))
        if name:
            out.append({'lvl': int(lvl), 'name': str(name).strip(), 'req': str(req or '-').strip(),
                        'max': int(ranks or 1), 'desc': str(desc or '').strip()})
    return sorted(out, key=lambda m: m['lvl'])


if __name__ == '__main__':
    wb = openpyxl.load_workbook(sys.argv[1])
    data = {'implants': implants(wb), 'merits': merits(wb)}
    sys.stdout.reconfigure(encoding='utf-8')
    print('// Generado por tools/build_catalog.py; no editar a mano.')
    print('globalThis.CATALOG=' + json.dumps(data, ensure_ascii=False, indent=0) + ';')
